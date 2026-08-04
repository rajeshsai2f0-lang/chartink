"""
╔══════════════════════════════════════════════════════════════╗
║   FINVIZ SCREENER AUTOMATION                                 ║
║   Runs multiple screeners → saves to Excel                   ║
╠══════════════════════════════════════════════════════════════╣
║  INSTALL ONCE (run in Command Prompt):                       ║
║   pip install requests beautifulsoup4 openpyxl pandas        ║
╚══════════════════════════════════════════════════════════════╝
"""

import sys, os, time, datetime

missing = []
for pkg in ['requests','bs4','openpyxl','pandas']:
    try: __import__(pkg)
    except ImportError: missing.append(pkg)
if missing:
    print(f"\n❌  Missing: {', '.join(missing)}")
    print(f"    Run:  pip install {' '.join(missing)}")
    if not os.environ.get("GITHUB_ACTIONS"):
        input("\nPress Enter to exit...")
    sys.exit(1)

import requests
from bs4 import BeautifulSoup
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ══════════════════════════════════════════════════════════════════════════════
#  ✏️  CREDENTIALS
#  Read from environment variables (FINVIZ_EMAIL / FINVIZ_PASSWORD) so nothing
#  sensitive is stored in this file. Set these as GitHub Actions secrets, or
#  locally via: set FINVIZ_EMAIL=... & set FINVIZ_PASSWORD=... (Windows)
# ══════════════════════════════════════════════════════════════════════════════
CREDENTIALS = {
    "email":    os.environ.get("FINVIZ_EMAIL", ""),
    "password": os.environ.get("FINVIZ_PASSWORD", ""),
}

if not CREDENTIALS["email"] or not CREDENTIALS["password"]:
    print("\n❌  Missing Finviz credentials.")
    print("    Set FINVIZ_EMAIL and FINVIZ_PASSWORD as environment variables")
    print("    (GitHub Actions secrets, or set them locally before running).")
    if not os.environ.get("GITHUB_ACTIONS"):
        input("\nPress Enter to exit...")
    sys.exit(1)

# ══════════════════════════════════════════════════════════════════════════════
#  ✏️  SCREENER URLs  (single clean list — no duplicates)
# ══════════════════════════════════════════════════════════════════════════════
SCREENERS = [
    (
        "Post-Market CANSLIM",
        "https://finviz.com/screener.ashx?v=111&p=d&f=cap_midover,fa_salesqoq_high,fa_salesyoyttm_high,ind_aerospacedefense|advertisingagencies|agriculturalinputs|airportsairservices|assetmanagement|apparelmanufacturing|airlines|aluminum|apparelretail|automanufacturers|autotruckdealerships|banksregional|beveragesnonalcoholic|buildingmaterials|businessequipmentsupplies|chemicals|closedendfundequity|autoparts|banksdiversified|beveragesbrewers|beverageswineriesdistilleries|broadcasting|buildingproductsequipment|capitalmarkets|closedendfunddebt|closedendfundforeign|cokingcoal|computerhardware|communicationequipment|conglomerates|consumerelectronics|creditservices|diagnosticsresearch|drugmanufacturersgeneral|educationtrainingservices|electroniccomponents|electronicscomputerdistribution|entertainment|farmheavyconstructionmachinery|financialconglomerates|fooddistribution|furnishingsfixturesappliances|gold|gambling|footwearaccessories|financialdatastockexchanges|farmproducts|exchangetradedfund|engineeringconstruction|electronicgamingmultimedia|electricalequipmentparts|discountstores|departmentstores|copper|consultingservices|confectioners|drugmanufacturersspecialtygeneric|healthcareplans|homeimprovementretail|industrialdistribution|infrastructureoperations|insurancediversified|insurancespecialty|insurancepropertycasualty|internetcontentinformation|leisure|lumberwoodproduction|marineshipping|medicaldevices|medicalinstrumentssupplies|mortgagefinance|oilgasep|grocerystores|healthinformationservices|householdpersonalproducts|informationtechnologyservices|insurancebrokers|insurancelife|insurancereinsurance|integratedfreightlogistics|internetretail|lodging|luxurygoods|medicalcarefacilities|medicaldistribution|metalfabrication|oilgasdrilling|oilgasequipmentservices|oilgasintegrated|oilgasrefiningmarketing|otherpreciousmetalsmining|otherindustrialmetalsmining|oilgasmidstream|packagedfoods|packagingcontainers|personalservices|paperpaperproducts|pharmaceuticalretailers|pollutiontreatmentcontrols|railroads|realestatediversified|recreationalvehicles|reithealthcarefacilities|reitindustrial|reitoffice|reitretail|rentalleasingservices|resortscasinos|scientifictechnicalinstruments|semiconductorequipmentmaterials|shellcompanies|softwareapplication|solar|specialtychemicals|specialtyretail|steel|textilemanufacturing|tobacco|travelservices|uranium|utilitiesindependentpowerproducers|utilitiesregulatedgas|utilitiesrenewable|publishing|realestatedevelopment|realestateservices|reitdiversified|reithotelmotel|reitmortgage|reitresidential|reitspecialty|residentialconstruction|restaurants|securityprotectionservices|semiconductors|silver|specialtyindustrialmachinery|softwareinfrastructure|specialtybusinessservices|staffingemploymentservices|telecomservices|thermalcoal|toolsaccessories|trucking|utilitiesdiversified|utilitiesregulatedwater|utilitiesregulatedelectric|wastemanagement,sh_avgvol_o2000,sh_curvol_o1000,sh_insttrans_pos,ta_highlow20d_a5h,ta_highlow50d_a5h,ta_volatility_wo4&ft=4&ta=0",
    ),
    (
        "Jeff Rocket",
        "https://finviz.com/screener.ashx?v=131&f=cap_smallover,ind_stocksonly,sh_avgvol_o1000,sh_float_u100,sh_short_o30&ft=4",
    ),
    (
        "Jeff IPO",
        "https://finviz.com/screener.ashx?v=211&f=cap_midover,fa_epsyoy1_pos,ipodate_prevyear,sh_avgvol_o1000&ft=4&o=industry",
    ),
    (
        "6 Month Exceeding 100",
        "https://finviz.com/screener.ashx?v=111&f=cap_smallover,sh_avgvol_o300,sh_curvol_o100,ta_perf_26w100o,ta_volatility_mo5&ft=4&o=-marketcap",
    ),
    (
        "3 Month Exceeding 50",
        "https://finviz.com/screener.ashx?v=111&f=cap_smallover,sh_avgvol_o300,sh_curvol_o100,ta_perf_13w50o,ta_volatility_mo5&ft=4&o=-marketcap",
    ),
    (
        "1 Month Exceeding 50",
        "https://finviz.com/screener.ashx?v=111&f=cap_smallover,sh_avgvol_o300,sh_curvol_o100,ta_perf_4w50o,ta_volatility_mo5&ft=4&o=-marketcap",
    ),
    (
        "1 Month Exceeding 30",
        "https://finviz.com/screener.ashx?v=111&f=cap_smallover,sh_avgvol_o300,sh_curvol_o100,ta_perf_4w30o,ta_volatility_mo5&ft=4&o=-marketcap",
    ),
    (
        "1 Week Exceeding 20",
        "https://finviz.com/screener.ashx?v=111&f=cap_smallover,sh_avgvol_o300,sh_curvol_o100,ta_perf_1w20o,ta_volatility_wo4&ft=4&o=-marketcap",
    ),
    (
        "TA Relative Volume",
        "https://finviz.com/screener.ashx?v=111&f=cap_midover,geo_usa,sh_avgvol_o500,sh_curvol_o1000,sh_price_o5,sh_relvol_o1,ta_sma20_pa,ta_sma200_pa,ta_sma50_pa&ft=3",
    ),
    (
        "Last Year IPO",
        "https://finviz.com/screener.ashx?v=111&f=ind_stocksonly,ipodate_prevyear&o=-change",
    ),
    (
        "Golden Stack Uptrend",
        "https://finviz.com/screener.ashx?v=211&p=w&f=geo_usa,ind_stocksonly,sh_avgvol_o1000,sh_relvol_o1,ta_averagetruerange_o1,ta_sma20_pa,ta_sma200_sb50,ta_sma50_sb20&ft=4&ta=0&o=-change&r=25",
    ),
    (
        "Breakout Screener",
        "https://finviz.com/screener.ashx?v=111&f=fa_ltdebteq_u1,fa_roe_o20,geo_usa,sh_avgvol_o100,ta_highlow50d_nh,ta_sma20_pa,ta_sma200_pa,ta_sma50_pa&ft=4",
    ),
    (
        "Post-Market Bases at Beaten Down Levels",
        "https://finviz.com/screener?v=111&p=d&f=cap_smallover,ind_advertisingagencies|agriculturalinputs|airportsairservices|apparelmanufacturing|assetmanagement|autoparts|banksdiversified|beveragesbrewers|wastemanagement|utilitiesregulatedwater|utilitiesregulatedelectric|utilitiesdiversified|trucking|toolsaccessories|thermalcoal|telecomservices|staffingemploymentservices|specialtyindustrialmachinery|specialtybusinessservices|softwareinfrastructure|silver|semiconductors|securityprotectionservices|utilitiesrenewable|utilitiesregulatedgas|utilitiesindependentpowerproducers|uranium|travelservices|tobacco|textilemanufacturing|steel|specialtyretail|specialtychemicals|solar|softwareapplication|shellcompanies|semiconductorequipmentmaterials|scientifictechnicalinstruments|restaurants|residentialconstruction|reitspecialty|reitresidential|reitmortgage|reithotelmotel|reitdiversified|realestateservices|realestatedevelopment|publishing|railroads|realestatediversified|recreationalvehicles|reithealthcarefacilities|reitindustrial|reitoffice|reitretail|rentalleasingservices|resortscasinos|pollutiontreatmentcontrols|personalservices|packagingcontainers|otherpreciousmetalsmining|oilgasrefiningmarketing|oilgasintegrated|oilgasep|mortgagefinance|medicalinstrumentssupplies|medicaldevices|marineshipping|lumberwoodproduction|leisure|internetcontentinformation|insurancespecialty|insurancepropertycasualty|insurancediversified|householdpersonalproducts|informationtechnologyservices|insurancebrokers|insurancelife|insurancereinsurance|integratedfreightlogistics|internetretail|lodging|luxurygoods|medicalcarefacilities|medicaldistribution|metalfabrication|oilgasdrilling|oilgasequipmentservices|oilgasmidstream|otherindustrialmetalsmining|packagedfoods|paperpaperproducts|pharmaceuticalretailers|infrastructureoperations|industrialdistribution|homeimprovementretail|healthcareplans|gold|financialconglomerates|farmheavyconstructionmachinery|entertainment|electronicscomputerdistribution|electroniccomponents|educationtrainingservices|discountstores|drugmanufacturersspecialtygeneric|electricalequipmentparts|electronicgamingmultimedia|engineeringconstruction|exchangetradedfund|farmproducts|financialdatastockexchanges|footwearaccessories|gambling|grocerystores|furnishingsfixturesappliances|fooddistribution|healthinformationservices|buildingmaterials|businessequipmentsupplies|chemicals|closedendfundequity|cokingcoal|computerhardware|conglomerates|consumerelectronics|creditservices|diagnosticsresearch|drugmanufacturersgeneral|departmentstores|copper|consultingservices|confectioners|communicationequipment|closedendfundforeign|closedendfunddebt|capitalmarkets|buildingproductsequipment|beverageswineriesdistilleries|beveragesnonalcoholic|broadcasting|banksregional|autotruckdealerships|automanufacturers|aluminum|apparelretail|airlines|aerospacedefense,sh_avgvol_o1000,sh_curvol_o1000,sh_insttrans_pos,sh_price_o1,ta_alltime_b70h,ta_highlow50d_a15h,ta_highlow52w_b30h,ta_perf_ytddown,ta_sma200_-20to20-a,ta_volatility_wo4&ft=4&ta=0&dr=y1",
    ),
    (
        "Hottest Screener",
        "https://finviz.com/screener.ashx?v=111&p=d&f=cap_0.15to,sh_avgvol_o2000,sh_curvol_o1000,sh_float_to500x,sh_insttrans_pos,sh_short_high,ta_perf_13w30o,ta_volatility_wo5&ft=4&ta=0&o=-industry",
    ),
    (
        "High Short Float",
        "https://finviz.com/screener.ashx?v=131&f=cap_smallover,sh_avgvol_o1000,sh_float_u100,sh_short_o30&ft=4",
    ),

(
        "Post-Market Hottest Stock Screener (I do daily)",
        "https://finviz.com/screener.ashx?vv=411&p=d&f=cap_0.15to,ind_advertisingagencies|agriculturalinputs|airportsairservices|apparelmanufacturing|assetmanagement|banksdiversified|aerospacedefense|airlines|aluminum|apparelretail|automanufacturers|autotruckdealerships|banksregional|beveragesnonalcoholic|buildingmaterials|autoparts|beveragesbrewers|beverageswineriesdistilleries|broadcasting|buildingproductsequipment|capitalmarkets|communicationequipment|closedendfunddebt|closedendfundforeign|businessequipmentsupplies|chemicals|closedendfundequity|computerhardware|cokingcoal|conglomerates|consumerelectronics|creditservices|diagnosticsresearch|drugmanufacturersgeneral|confectioners|consultingservices|copper|departmentstores|discountstores|educationtrainingservices|electroniccomponents|electronicscomputerdistribution|entertainment|farmheavyconstructionmachinery|drugmanufacturersspecialtygeneric|electricalequipmentparts|electronicgamingmultimedia|engineeringconstruction|exchangetradedfund|financialconglomerates|fooddistribution|furnishingsfixturesappliances|gold|healthcareplans|farmproducts|financialdatastockexchanges|footwearaccessories|gambling|grocerystores|healthinformationservices|homeimprovementretail|industrialdistribution|householdpersonalproducts|informationtechnologyservices|infrastructureoperations|insurancediversified|insurancepropertycasualty|insurancespecialty|internetcontentinformation|leisure|lumberwoodproduction|insurancebrokers|insurancelife|insurancereinsurance|integratedfreightlogistics|internetretail|lodging|marineshipping|medicaldevices|medicalinstrumentssupplies|mortgagefinance|oilgasep|oilgasintegrated|luxurygoods|medicalcarefacilities|medicaldistribution|oilgasdrilling|oilgasequipmentservices|metalfabrication|oilgasrefiningmarketing|otherpreciousmetalsmining|packagingcontainers|personalservices|pollutiontreatmentcontrols|railroads|oilgasmidstream|otherindustrialmetalsmining|packagedfoods|paperpaperproducts|pharmaceuticalretailers|publishing|realestatediversified|realestatedevelopment|realestateservices|reitdiversified|reithotelmotel|reitmortgage|reitresidential|recreationalvehicles|reithealthcarefacilities|reitindustrial|reitoffice|reitretail|rentalleasingservices|reitspecialty|resortscasinos|scientifictechnicalinstruments|semiconductorequipmentmaterials|shellcompanies|softwareapplication|solar|specialtychemicals|specialtyretail|steel|textilemanufacturing|tobacco|travelservices|uranium|utilitiesindependentpowerproducers|utilitiesregulatedgas|utilitiesrenewable|residentialconstruction|restaurants|securityprotectionservices|semiconductors|softwareinfrastructure|silver|specialtybusinessservices|staffingemploymentservices|specialtyindustrialmachinery|telecomservices|thermalcoal|toolsaccessories|trucking|utilitiesdiversified|utilitiesregulatedelectric|utilitiesregulatedwater|wastemanagement,sh_avgvol_o2000,sh_curvol_o1000,sh_float_to500x,sh_insttrans_pos,sh_short_high,ta_perf_13w30o,ta_volatility_wo5&ft=4&ta=0&o=-industry",
    ),


   (
        "praddep9m",
        "https://finviz.com/screener.ashx?v=111&f=cap_smallover,ind_stocksonly,sh_curvol_o10000,sh_relvol_o1,ta_change_u,ta_changeopen_u,ta_perf_1wup,ta_perf2_4wup&ft=4&o=-marketcap",
    ),

     (
        "Post-Market Highest Short Float (I do it daily)",
        "https://finviz.com/screener.ashx?v=131&f=cap_smallover,ind_advertisingagencies|airportsairservices|agriculturalinputs|aerospacedefense|aluminum|airlines|apparelretail|automanufacturers|autotruckdealerships|beveragesnonalcoholic|banksregional|buildingmaterials|businessequipmentsupplies|broadcasting|beverageswineriesdistilleries|banksdiversified|apparelmanufacturing|assetmanagement|autoparts|beveragesbrewers|buildingproductsequipment|capitalmarkets|closedendfunddebt|closedendfundforeign|communicationequipment|confectioners|consultingservices|copper|departmentstores|discountstores|drugmanufacturersspecialtygeneric|electricalequipmentparts|electronicgamingmultimedia|engineeringconstruction|exchangetradedfund|farmproducts|financialdatastockexchanges|footwearaccessories|gambling|grocerystores|healthinformationservices|householdpersonalproducts|informationtechnologyservices|chemicals|closedendfundequity|cokingcoal|computerhardware|conglomerates|consumerelectronics|creditservices|diagnosticsresearch|drugmanufacturersgeneral|educationtrainingservices|electroniccomponents|electronicscomputerdistribution|entertainment|farmheavyconstructionmachinery|financialconglomerates|fooddistribution|furnishingsfixturesappliances|gold|healthcareplans|homeimprovementretail|industrialdistribution|infrastructureoperations|insurancediversified|insurancepropertycasualty|insurancespecialty|internetcontentinformation|utilitiesrenewable|utilitiesregulatedgas|utilitiesindependentpowerproducers|uranium|travelservices|tobacco|textilemanufacturing|steel|specialtyretail|specialtychemicals|solar|softwareapplication|shellcompanies|semiconductorequipmentmaterials|scientifictechnicalinstruments|resortscasinos|rentalleasingservices|reitretail|reitoffice|reitindustrial|reithealthcarefacilities|recreationalvehicles|realestatediversified|railroads|pollutiontreatmentcontrols|personalservices|packagingcontainers|otherpreciousmetalsmining|oilgasrefiningmarketing|oilgasintegrated|oilgasep|mortgagefinance|medicalinstrumentssupplies|medicaldevices|marineshipping|lumberwoodproduction|leisure|insurancebrokers|insurancelife|insurancereinsurance|integratedfreightlogistics|internetretail|lodging|luxurygoods|medicalcarefacilities|medicaldistribution|metalfabrication|oilgasdrilling|oilgasequipmentservices|oilgasmidstream|otherindustrialmetalsmining|packagedfoods|paperpaperproducts|pharmaceuticalretailers|publishing|realestatedevelopment|realestateservices|reitdiversified|reithotelmotel|reitmortgage|reitresidential|reitspecialty|residentialconstruction|restaurants|securityprotectionservices|semiconductors|silver|softwareinfrastructure|specialtybusinessservices|specialtyindustrialmachinery|staffingemploymentservices|telecomservices|thermalcoal|toolsaccessories|trucking|utilitiesdiversified|utilitiesregulatedelectric|utilitiesregulatedwater|wastemanagement,sh_avgvol_o1000,sh_float_u100,sh_short_o30&ft=4",
    ),
]

OUTPUT_FILE = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    f"Finviz_Screener_{datetime.date.today().strftime('%Y-%m-%d')}.xlsx"
)

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/122.0.0.0 Safari/537.36"
    ),
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "en-US,en;q=0.9",
    "Referer": "https://finviz.com/",
}

NAV_LABELS = {
    "overview", "valuation", "financial", "ownership", "performance", "technical",
    "custom", "charts", "tickers", "basic", "ta", "news", "snapshot", "maps", "stats",
    "etf", "etf perf", "#", "no.", "",
}

TICKERS_PER_ROW = 95   # ← change this if you ever need a different batch size
                        #    (still used for the per-screener tabs)

def trim_first_letter(sym):
    """
    Finviz is currently prefixing an extra letter onto every symbol
    (e.g. 'AABNB' instead of 'ABNB'). Strip the first character so we
    get back the real ticker. Leaves 1-character symbols untouched.
    """
    s = str(sym).strip()
    return s[1:] if len(s) > 1 else s

# ══════════════════════════════════════════════════════════════════════════════
#  LOGIN
# ══════════════════════════════════════════════════════════════════════════════
def login(session):
    print("\n🔐  Logging into Finviz...")
    try:
        session.get("https://finviz.com/login.ashx", headers=HEADERS, timeout=15)
        time.sleep(1.5)
        resp = session.post(
            "https://finviz.com/login-email?remember=true",
            data={"email": CREDENTIALS["email"], "password": CREDENTIALS["password"], "remember": "1"},
            headers={**HEADERS, "Referer": "https://finviz.com/login.ashx",
                     "Content-Type": "application/x-www-form-urlencoded"},
            timeout=15, allow_redirects=True,
        )
        if "logout" in resp.text.lower() or "Sign Out" in resp.text:
            print("   ✅  Logged in")
            return True
        print("   ⚠️  Login uncertain — continuing (Elite needed for all columns)")
        return False
    except Exception as e:
        print(f"   ❌  {e}")
        return False

# ══════════════════════════════════════════════════════════════════════════════
#  SCRAPE
# ══════════════════════════════════════════════════════════════════════════════
def _find_results_table(soup):
    for t in soup.find_all("table"):
        if t.find("a", class_="screener-link-primary"):
            return t, "strategy-1:screener-link-primary"
    for t in soup.find_all("table"):
        if t.find("td", class_=lambda c: c and "screener-body-table" in c):
            return t, "strategy-2:screener-body-table"
    for t in soup.find_all("table"):
        if t.find("th", class_=lambda c: c and "table-top" in c):
            return t, "strategy-3:table-top"
    best, best_cols, best_strat = None, 0, ""
    for i, t in enumerate(soup.find_all("table")):
        fr = t.find("tr")
        if fr:
            ncols = len(fr.find_all(["td", "th"]))
            if ncols > best_cols:
                best, best_cols, best_strat = t, ncols, f"strategy-4:widest({ncols}cols)"
    if best and best_cols >= 5:
        return best, best_strat
    return None, "not-found"


def scrape_screener(session, name, url):
    print(f"\n📡  Scraping: {name}")
    all_rows = []
    col_hdrs = None
    row_num  = 1

    while True:
        page_url = url.strip() if row_num == 1 else f"{url.strip()}&r={row_num}"
        print(f"   row {row_num}...", end=" ", flush=True)

        try:
            resp = session.get(page_url, headers=HEADERS, timeout=25)

            if "login.ashx" in resp.url:
                print("⚠️  Redirected to login — check credentials")
                break

            soup = BeautifulSoup(resp.text, "html.parser")

            page_text = soup.get_text()
            if "No results found" in page_text or "0 Total" in page_text:
                print("(0 results)")
                break

            table, strat = _find_results_table(soup)

            if not table:
                print(f"\n   ❌  Table not found.")
                print(f"       URL: {resp.url}")
                print(f"       Title: {soup.title.string if soup.title else 'N/A'}")
                all_tables = soup.find_all("table")
                print(f"       Tables on page: {len(all_tables)}")
                for i, t in enumerate(all_tables[:5]):
                    fr = t.find("tr")
                    ncols = len(fr.find_all(["td","th"])) if fr else 0
                    print(f"       Table {i}: {ncols} cols, class={t.get('class','')}")
                break

            print(f"[{strat}]", end=" ", flush=True)
            rows = table.find_all("tr")

            if col_hdrs is None:
                first_row = rows[0] if rows else None
                if first_row:
                    cells = first_row.find_all("th") or first_row.find_all("td")
                    col_hdrs = [c.get_text(strip=True) or f"Col{i}" for i, c in enumerate(cells)]
                    print(f"\n   Columns ({len(col_hdrs)}): {col_hdrs[:8]}...")
                if not col_hdrs:
                    print("   ❌  Could not extract headers")
                    break

            data_rows = []
            for row in rows[1:]:
                cells     = row.find_all("td")
                cells_txt = [c.get_text(strip=True) for c in cells]
                if not cells_txt or all(v == "" for v in cells_txt):
                    continue
                if cells_txt[0].strip().lower() in NAV_LABELS:
                    continue
                if len(cells_txt) != len(col_hdrs):
                    continue
                data_rows.append(cells_txt)

            if not data_rows:
                print("✅  No more data")
                break

            all_rows.extend(data_rows)
            print(f"✅  {len(data_rows)} rows")

            if len(data_rows) < 20:
                break

            row_num += len(data_rows)
            time.sleep(1.5)

        except Exception as e:
            print(f"❌  {e}")
            import traceback; traceback.print_exc()
            break

    if not all_rows:
        print(f"   ⚠️  No data for: {name}")
        return pd.DataFrame()

    df = pd.DataFrame(all_rows, columns=col_hdrs)
    df = df.dropna(how="all").drop_duplicates()
    print(f"   📋  {len(df)} tickers total")
    return df

# ══════════════════════════════════════════════════════════════════════════════
#  TICKER COLUMN DETECTION
# ══════════════════════════════════════════════════════════════════════════════
def find_ticker_col(df):
    for col in df.columns:
        if col.strip().lower() in ("ticker", "symbol", "t"):
            return col
    for col in df.columns:
        s = df[col].dropna().astype(str).head(15)
        if s.str.match(r"^[A-Z]{1,5}$").mean() > 0.5:
            return col
    return df.columns[1] if len(df.columns) > 1 else df.columns[0]


# ══════════════════════════════════════════════════════════════════════════════
#  HELPER — write ticker batch rows onto a worksheet
#  Returns the next available row number after the block.
# ══════════════════════════════════════════════════════════════════════════════
def write_ticker_batches(ws, tickers, start_row, n_cols,
                         fill_fn, thin_fn, medium_fn,
                         C_DARK, C_WHITE, C_NOTE, C_HDR, C_GOLD):
    """
    Writes a labelled header + one merged cell per batch of TICKERS_PER_ROW
    tickers onto *ws*, starting at *start_row*.
    Returns the row number of the first empty row after the block.
    """
    last_col = get_column_letter(n_cols)
    bd_side  = Side(style="medium", color=C_HDR)
    bd       = Border(left=bd_side, right=bd_side, top=bd_side, bottom=bd_side)

    # ── section label ─────────────────────────────────────────────────────────
    ws.row_dimensions[start_row].height = 18
    ws.merge_cells(f"A{start_row}:{last_col}{start_row}")
    c = ws[f"A{start_row}"]
    n_batches = (len(tickers) + TICKERS_PER_ROW - 1) // TICKERS_PER_ROW
    c.value     = (f"  ✂  COPY-PASTE TICKER LISTS — {len(tickers)} tickers   "
                   f"|   {n_batches} batch{'es' if n_batches != 1 else ''} of ≤{TICKERS_PER_ROW}")
    c.font      = Font(name="Arial", bold=True, size=10, color=C_WHITE)
    c.fill      = fill_fn(C_DARK)
    c.alignment = Alignment(horizontal="left", vertical="center")
    current_row = start_row + 1

    if not tickers:
        ws.row_dimensions[current_row].height = 18
        ws.merge_cells(f"A{current_row}:{last_col}{current_row}")
        c = ws[f"A{current_row}"]
        c.value     = "  (no tickers)"
        c.font      = Font(name="Arial", italic=True, size=9, color="888888")
        c.fill      = fill_fn(C_NOTE)
        c.alignment = Alignment(horizontal="left", vertical="center")
        return current_row + 2   # blank gap

    # ── one row per batch ─────────────────────────────────────────────────────
    for batch_idx in range(0, len(tickers), TICKERS_PER_ROW):
        batch      = tickers[batch_idx : batch_idx + TICKERS_PER_ROW]
        batch_num  = batch_idx // TICKERS_PER_ROW + 1
        batch_text = ", ".join(batch)

        # sub-label row
        ws.row_dimensions[current_row].height = 14
        ws.merge_cells(f"A{current_row}:{last_col}{current_row}")
        lbl = ws[f"A{current_row}"]
        lbl.value     = (f"  Batch {batch_num}  ({len(batch)} tickers  —  "
                         f"rows {batch_idx + 1}–{batch_idx + len(batch)})")
        lbl.font      = Font(name="Arial", bold=True, size=8, color=C_WHITE)
        lbl.fill      = fill_fn(C_HDR)
        lbl.alignment = Alignment(horizontal="left", vertical="center")
        current_row  += 1

        # ticker string row
        ws.row_dimensions[current_row].height = 55
        ws.merge_cells(f"A{current_row}:{last_col}{current_row}")
        c = ws[f"A{current_row}"]
        c.value     = batch_text
        c.font      = Font(name="Courier New", size=9)
        c.fill      = fill_fn(C_NOTE)
        c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        c.border    = bd
        current_row += 1

    return current_row + 1   # blank gap row after the block


# ══════════════════════════════════════════════════════════════════════════════
#  HELPER — write ONE consolidated comma-separated ticker list (no 95-batching)
#  Returns the next available row number after the block.
# ══════════════════════════════════════════════════════════════════════════════
def write_consolidated_list(ws, tickers, start_row, n_cols,
                             fill_fn, thin_fn,
                             C_DARK, C_WHITE, C_NOTE, C_HDR):
    """
    Writes ALL tickers as a single consolidated, comma-separated list
    (no more splitting into batches of TICKERS_PER_ROW).
    """
    last_col = get_column_letter(n_cols)
    bd_side  = Side(style="medium", color=C_HDR)
    bd       = Border(left=bd_side, right=bd_side, top=bd_side, bottom=bd_side)

    ws.row_dimensions[start_row].height = 18
    ws.merge_cells(f"A{start_row}:{last_col}{start_row}")
    c = ws[f"A{start_row}"]
    c.value     = f"  ✂  COPY-PASTE TICKER LIST — {len(tickers)} tickers (consolidated)"
    c.font      = Font(name="Arial", bold=True, size=10, color=C_WHITE)
    c.fill      = fill_fn(C_DARK)
    c.alignment = Alignment(horizontal="left", vertical="center")
    current_row = start_row + 1

    if not tickers:
        ws.row_dimensions[current_row].height = 18
        ws.merge_cells(f"A{current_row}:{last_col}{current_row}")
        c = ws[f"A{current_row}"]
        c.value     = "  (no tickers)"
        c.font      = Font(name="Arial", italic=True, size=9, color="888888")
        c.fill      = fill_fn(C_NOTE)
        c.alignment = Alignment(horizontal="left", vertical="center")
        return current_row + 2   # blank gap

    ticker_text = ", ".join(tickers)

    # Excel single-cell limit is 32,767 chars — scale row height with length
    # so the whole list stays readable/wrapped inside the one cell.
    row_height = 55 + min(len(ticker_text) // 60, 400)

    ws.row_dimensions[current_row].height = row_height
    ws.merge_cells(f"A{current_row}:{last_col}{current_row}")
    c = ws[f"A{current_row}"]
    c.value     = ticker_text
    c.font      = Font(name="Courier New", size=9)
    c.fill      = fill_fn(C_NOTE)
    c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    c.border    = bd
    current_row += 1

    return current_row + 1   # blank gap row after the block


# ══════════════════════════════════════════════════════════════════════════════
#  BUILD EXCEL
# ══════════════════════════════════════════════════════════════════════════════
def build_excel(screener_results, output_path):
    wb = Workbook()
    wb.remove(wb.active)

    def fill(h):
        return PatternFill("solid", fgColor=h)

    def thin(c="CCCCCC"):
        s = Side(style="thin", color=c)
        return Border(left=s, right=s, top=s, bottom=s)

    def medium(c="888888"):
        s = Side(style="medium", color=c)
        return Border(left=s, right=s, top=s, bottom=s)

    C_DARK  = "1A3A5C"; C_HDR  = "2E5FA3"; C_GOLD = "FFD700"; C_WHITE = "FFFFFF"
    C_S1    = "FFFFFF"; C_S2   = "EBF3FB"; C_TICK = "F0F7FF"
    C_POSB  = "E6F4EA"; C_POSF = "137333"; C_NEGB = "FCE8E6"; C_NEGF  = "C5221F"
    C_NOTE  = "FFFBEA"; C_GREEN = "00763D"; C_RED  = "C00000"

    now_str = datetime.datetime.now().strftime("%d %b %Y  %H:%M")
    all_tickers_flat = []
    tab_colors = ["2E75B6","70AD47","C00000","FF8C00","7030A0","00B0F0",
                  "BF9000","375623","4472C4","ED7D31","A9D18E","FF0000",
                  "0070C0","7F7F7F","FFC000"]

    for idx, (name, df) in enumerate(screener_results):
        safe = (name[:31].replace("/","_").replace("\\","_")
                .replace("?","").replace("*","")
                .replace("[","").replace("]","").replace(":",""))
        ws = wb.create_sheet(title=safe)
        ws.sheet_properties.tabColor = tab_colors[idx % len(tab_colors)]

        n_cols   = max(len(df.columns), 5) if not df.empty else 5
        last_col = get_column_letter(n_cols)

        # ── title banner ──────────────────────────────────────────────────────
        ws.row_dimensions[1].height = 5
        ws.row_dimensions[2].height = 28
        ws.row_dimensions[3].height = 14

        ws.merge_cells(f"A2:{last_col}2")
        c = ws["A2"]
        c.value     = f"  {name.upper()}   |   {len(df)} tickers   |   {now_str}"
        c.font      = Font(name="Arial", bold=True, size=12, color=C_WHITE)
        c.fill      = fill(C_DARK)
        c.alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells(f"A3:{last_col}3")
        c = ws["A3"]
        c.value     = f"  Source: finviz.com   |   Run: {now_str}"
        c.font      = Font(name="Arial", italic=True, size=8, color="AAAAAA")
        c.fill      = fill(C_DARK)
        c.alignment = Alignment(horizontal="left", vertical="center")

        if df.empty:
            ws["A5"].value = "⚠️  No data retrieved for this screener"
            ws["A5"].font  = Font(name="Arial", size=11, color=C_RED)
            continue

        pct_cols = {col for col in df.columns
                    if df[col].dropna().astype(str).head(20).str.contains(r"%").mean() > 0.4}

        ticker_col  = find_ticker_col(df)
        ticker_cidx = list(df.columns).index(ticker_col) + 1

        # Find Industry column and filter Shell Companies for ticker list
        industry_col = next(
            (col for col in df.columns if col.strip().lower() == "industry"), None
        )
        if industry_col:
            filtered_df = df[df[industry_col].astype(str).str.strip().str.lower() != "shell companies"]
            excluded = len(df) - len(filtered_df)
            if excluded:
                print(f"   🚫  Excluded {excluded} Shell Companies from summary")
        else:
            filtered_df = df

        sheet_tickers = (filtered_df[ticker_col].dropna().astype(str).str.strip()
                         .replace("", pd.NA).dropna().tolist())
        sheet_tickers = [t for t in sheet_tickers if t.upper() not in ("NAN", "#", "")]
        all_tickers_flat.extend(sheet_tickers)

        # ── ticker batch rows (inserted at row 4 onwards) ─────────────────────
        next_row = write_ticker_batches(
            ws, sheet_tickers, start_row=4,
            n_cols=n_cols,
            fill_fn=fill, thin_fn=thin, medium_fn=medium,
            C_DARK=C_DARK, C_WHITE=C_WHITE, C_NOTE=C_NOTE,
            C_HDR=C_HDR, C_GOLD=C_GOLD,
        )

        # blank spacer row
        ws.row_dimensions[next_row].height = 6
        next_row += 1

        # ── column headers for data table ─────────────────────────────────────
        hdr_row = next_row
        ws.row_dimensions[hdr_row].height = 22
        for ci, col_name in enumerate(df.columns, 1):
            c = ws.cell(row=hdr_row, column=ci)
            c.value     = col_name
            c.font      = Font(name="Arial", bold=True, size=9, color=C_WHITE)
            c.fill      = fill(C_HDR)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border    = medium()
            ws.column_dimensions[get_column_letter(ci)].width = max(len(str(col_name)) + 3, 9)

        # ── data rows ─────────────────────────────────────────────────────────
        for ri, (_, row) in enumerate(df.iterrows()):
            er = hdr_row + 1 + ri
            ws.row_dimensions[er].height = 15
            bg = C_S1 if ri % 2 == 0 else C_S2
            for ci, (col_name, val) in enumerate(zip(df.columns, row), 1):
                c   = ws.cell(row=er, column=ci)
                txt = str(val).strip() if pd.notna(val) else ""
                c.value = txt
                if ci == ticker_cidx:
                    c.font      = Font(name="Arial", bold=True, size=9, color=C_DARK)
                    c.fill      = fill(C_TICK)
                    c.alignment = Alignment(horizontal="center", vertical="center")
                    c.border    = thin("AACCEE")
                elif col_name in pct_cols and txt:
                    neg = txt.startswith("-")
                    c.font      = Font(name="Arial", size=9, color=C_NEGF if neg else C_POSF)
                    c.fill      = fill(C_NEGB if neg else C_POSB)
                    c.alignment = Alignment(horizontal="right", vertical="center")
                    c.border    = thin()
                else:
                    c.font      = Font(name="Arial", size=9)
                    c.fill      = fill(bg)
                    c.alignment = Alignment(horizontal="center", vertical="center")
                    c.border    = thin()

        ws.freeze_panes = f"A{hdr_row + 1}"
        ws.auto_filter.ref = (
            f"A{hdr_row}:{last_col}{hdr_row + len(df)}"
        )

    # ── Summary tabs (deduplicated ticker lists) ────────────────────────────────
    from collections import Counter

    # Raw (as-scraped, untouched) — dedup + frequency
    counts_raw = Counter(t.upper() for t in all_tickers_flat)
    seen, unique_raw = set(), []
    for t in all_tickers_flat:
        u = t.upper()
        if u not in seen:
            seen.add(u); unique_raw.append(u)
    unique_raw.sort()

    # Trimmed (first alphabet stripped — see trim_first_letter()) — dedup + frequency
    trimmed_flat = [trim_first_letter(t) for t in all_tickers_flat]
    counts_trim = Counter(t.upper() for t in trimmed_flat)
    seen, unique_trim = set(), []
    for t in trimmed_flat:
        u = t.upper()
        if u not in seen:
            seen.add(u); unique_trim.append(u)
    unique_trim.sort()

    # ── Tab 1: "All Tickers" — trimmed symbols, ONE consolidated list ──────────
    ws_s = wb.create_sheet(title="All Tickers")
    ws_s.sheet_properties.tabColor = "FFD700"
    ws_s.row_dimensions[1].height  = 5
    ws_s.row_dimensions[2].height  = 30
    ws_s.row_dimensions[3].height  = 14

    ws_s.merge_cells("A2:G2")
    c = ws_s["A2"]
    c.value     = f"  ALL TICKERS — TRIMMED & DEDUPLICATED   |   {len(unique_trim)} unique   |   {now_str}"
    c.font      = Font(name="Arial", bold=True, size=13, color="000000")
    c.fill      = fill(C_GOLD)
    c.alignment = Alignment(horizontal="left", vertical="center")

    ws_s.merge_cells("A3:G3")
    c = ws_s["A3"]
    c.value     = (f"  First alphabet trimmed from each symbol   |   "
                   f"Total raw: {len(all_tickers_flat)}   |   After dedup: {len(unique_trim)}")
    c.font      = Font(name="Arial", italic=True, size=9, color="555555")
    c.fill      = fill(C_GOLD)
    c.alignment = Alignment(horizontal="left", vertical="center")

    for cl, w in [("A",6),("B",14),("C",20),("D",12),("E",12),("F",12),("G",12)]:
        ws_s.column_dimensions[cl].width = w

    next_row_s = write_consolidated_list(
        ws_s, unique_trim, start_row=5,
        n_cols=7,
        fill_fn=fill, thin_fn=thin,
        C_DARK=C_DARK, C_WHITE=C_WHITE, C_NOTE=C_NOTE, C_HDR=C_HDR,
    )

    # ── per-ticker frequency table (trimmed) ────────────────────────────────────
    tbl_start = next_row_s + 1
    ws_s.row_dimensions[tbl_start].height = 20
    for cl, hdr in [("A","#"),("B","Ticker"),("C","In # Screeners"),("D",""),("E",""),("F",""),("G","")]:
        c = ws_s[f"{cl}{tbl_start}"]
        c.value     = hdr
        c.font      = Font(name="Arial", bold=True, size=9, color=C_WHITE)
        c.fill      = fill(C_HDR)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border    = medium()

    for i, ticker in enumerate(unique_trim):
        r   = tbl_start + 1 + i
        bg  = C_S1 if i % 2 == 0 else C_S2
        cnt = counts_trim.get(ticker, 1)
        ws_s.row_dimensions[r].height = 15
        for cl, val in [("A", i + 1), ("B", ticker), ("C", cnt)]:
            c = ws_s[f"{cl}{r}"]
            c.value     = val
            c.font      = Font(name="Arial", size=9, bold=(cl == "B"),
                               color=C_GREEN if cnt > 1 else "000000")
            c.fill      = fill(C_S2 if cnt > 1 else bg)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border    = thin()

    ws_s.freeze_panes = f"A{tbl_start + 1}"

    # ── Tab 2: "All Tickers Raw" — untouched symbols, ONE consolidated list ────
    ws_r = wb.create_sheet(title="All Tickers Raw")
    ws_r.sheet_properties.tabColor = "BF9000"
    ws_r.row_dimensions[1].height  = 5
    ws_r.row_dimensions[2].height  = 30
    ws_r.row_dimensions[3].height  = 14

    ws_r.merge_cells("A2:G2")
    c = ws_r["A2"]
    c.value     = f"  ALL TICKERS — RAW (AS SCRAPED, NOT TRIMMED)   |   {len(unique_raw)} unique   |   {now_str}"
    c.font      = Font(name="Arial", bold=True, size=13, color="000000")
    c.fill      = fill(C_GOLD)
    c.alignment = Alignment(horizontal="left", vertical="center")

    ws_r.merge_cells("A3:G3")
    c = ws_r["A3"]
    c.value     = (f"  No trimming applied   |   "
                   f"Total raw: {len(all_tickers_flat)}   |   After dedup: {len(unique_raw)}")
    c.font      = Font(name="Arial", italic=True, size=9, color="555555")
    c.fill      = fill(C_GOLD)
    c.alignment = Alignment(horizontal="left", vertical="center")

    for cl, w in [("A",6),("B",14),("C",20),("D",12),("E",12),("F",12),("G",12)]:
        ws_r.column_dimensions[cl].width = w

    write_consolidated_list(
        ws_r, unique_raw, start_row=5,
        n_cols=7,
        fill_fn=fill, thin_fn=thin,
        C_DARK=C_DARK, C_WHITE=C_WHITE, C_NOTE=C_NOTE, C_HDR=C_HDR,
    )

    # ── sheet order: "All Tickers" (trimmed) first, "All Tickers Raw" second ───
    def _move_to_front(name):
        current = wb.sheetnames.index(name)
        wb.move_sheet(name, offset=-current)

    _move_to_front("All Tickers Raw")
    _move_to_front("All Tickers")

    wb.save(output_path)
    print(f"\n✅  Excel saved: {output_path}")
    return unique_trim

# ══════════════════════════════════════════════════════════════════════════════
#  MAIN
# ══════════════════════════════════════════════════════════════════════════════
if __name__ == "__main__":
    print("=" * 60)
    print("  FINVIZ SCREENER AUTOMATION")
    print("=" * 60)

    session = requests.Session()
    login(session)

    results = []
    for name, url in SCREENERS:
        df = scrape_screener(session, name, url)
        results.append((name, df))
        time.sleep(2.5)

    if not results:
        print("\n❌  No data. Check URLs and credentials.")
        if not os.environ.get("GITHUB_ACTIONS"):
            input("\nPress Enter to exit...")
        sys.exit(1)

    print(f"\n\n📊  Building Excel...")
    unique = build_excel(results, OUTPUT_FILE)
    print(f"\n🎯  {len(unique)} unique tickers")
    print(f"    Preview: {', '.join(unique[:10])}{'...' if len(unique) > 10 else ''}")

    if not os.environ.get("GITHUB_ACTIONS"):
        try:
            os.startfile(OUTPUT_FILE)
        except Exception:
            import subprocess
            subprocess.Popen(["start", OUTPUT_FILE], shell=True)

        input("\nPress Enter to close...")
    else:
        print(f"::notice::Excel file ready at {OUTPUT_FILE}")
